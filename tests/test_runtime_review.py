"""관리 진입·원천 표·알림 결과의 독립 합성 연결 검토."""
import builtins
import copy
import io
import json
from pathlib import Path
import runpy
import sys

import openpyxl
import pytest
import requests

import monitor
import position_cli
import stop_manager
import telegram_bot
from test_portfolio_configuration import load_config
from test_silent_failures import env, _market


@pytest.mark.parametrize("cell_error", ["#DIV/0!", "#VALUE!", "#N/A"])
def test_excel_error_is_not_an_optional_blank_entry_price(load_config, cell_error, monkeypatch):
    module, _ = load_config()
    original = ({"SYNTH-OLD": "synthetic-name"}, {"SYNTH-OLD": ["synthetic-account"]}, {"SYNTH-OLD": 5.})
    module._kr_names_from_drive, module._symbol_accounts, module._symbol_entry_prices = copy.deepcopy(original)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Ticker", "구분", "진입가격"])
    sheet.append(["SYNTH-GOOD", "미국", "10"])
    sheet.append(["SYNTH-BAD", "미국", cell_error])
    assert sheet["C3"].data_type == "e"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    loader = openpyxl.load_workbook
    closed = []
    def tracked_loader(*args, **kwargs):
        book = loader(*args, **kwargs)
        close = book.close
        def tracked_close():
            closed.append(True)
            return close()
        book.close = tracked_close
        return book
    monkeypatch.setattr(openpyxl, "load_workbook", tracked_loader)
    with pytest.raises(ValueError, match="^portfolio_"):
        module._parse_portfolio_df(module._portfolio_frame_from_bytes(output.getvalue(), excel=True))
    assert closed == [True]
    assert (module._kr_names_from_drive, module._symbol_accounts, module._symbol_entry_prices) == original


@pytest.mark.parametrize("args,code", [(["--add-pos", "SYNTH", "bad", "2"], 1), (["--add-pos", "--help"], 0)])
def test_actual_administration_exits_before_operational_imports(monkeypatch, args, code):
    monkeypatch.setenv("GITHUB_ACTIONS", "false")
    monkeypatch.setattr(sys, "argv", ["monitor.py", *args])
    original = builtins.__import__
    def guarded(name, *values, **kwargs):
        if name in {"config", "data_collector", "telegram_bot", "visualizer", "schedule", "log_masking"}:
            pytest.fail("administrative parser reached an operational import")
        return original(name, *values, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    with pytest.raises(SystemExit) as result:
        runpy.run_path(str(Path(monitor.__file__)), run_name="__main__")
    assert result.value.code == code


@pytest.mark.parametrize("ack", [False, True])
def test_one_shot_receipt_or_local_storage_failure_never_reports_success(env, monkeypatch, ack, caplog):
    symbol = "SYNTH-REVIEW"
    stop_manager.add_position(symbol, 80., 70., 100.)
    old = env.state.read_bytes()
    _market(env, [symbol], {symbol: ["SURGE DOWN -6%"]})
    monkeypatch.setattr(monitor, "STATE_FILE", env.state)
    monkeypatch.setattr(monitor, "ALL_SYMBOLS", [symbol])
    monkeypatch.setattr(monitor, "IS_GITHUB_ACTIONS", False)
    monkeypatch.setattr(monitor._config, "PORTFOLIO_ERROR", "")
    monkeypatch.setattr(sys, "argv", ["monitor.py", "--stop-check"])
    monkeypatch.setattr(telegram_bot, "_is_configured", lambda: True)
    response = requests.Response()
    response.status_code = 200
    response._content = json.dumps({"ok": ack}).encode()
    calls = []
    def post(*args, **kwargs):
        calls.append("post")
        return response
    monkeypatch.setattr(telegram_bot.requests, "post", post)
    monkeypatch.setattr(env.tg, "send_message", telegram_bot.send_message)
    if ack:
        def write_failed(*args, **kwargs):
            raise stop_manager.StateValidationError("synthetic_write_failed")
        monkeypatch.setattr(stop_manager, "write_state", write_failed)
    with pytest.raises(SystemExit) as result:
        monitor.main()
    assert result.value.code == 1 and monitor._problems
    assert calls and env.state.read_bytes() == old
    assert json.loads(old)["alert_log"] == {}


@pytest.mark.parametrize("status", [429, 503])
def test_invalid_negative_retry_delay_is_a_sanitized_failure(monkeypatch, status):
    monkeypatch.setattr(telegram_bot, "_is_configured", lambda: True)
    response = requests.Response()
    response.status_code = status
    response._content = b'{"ok":false,"parameters":{"retry_after":-1}}'
    calls, delays = [], []
    def post(*args, **kwargs):
        calls.append("post")
        return response
    def sleep(delay):
        # 실제 sleep(-1)의 ValueError를 시간 대기 없이 재현한다.
        if delay < 0:
            raise ValueError("sleep length must be non-negative")
        delays.append(delay)
    monkeypatch.setattr(telegram_bot.requests, "post", post)
    monkeypatch.setattr(telegram_bot.time, "sleep", sleep)
    assert telegram_bot.send_message("synthetic") is False
    assert len(calls) == telegram_bot.SEND_ATTEMPTS == 3
    assert delays == [2, 4, 8]


def test_gha_administration_script_is_rejected_before_operational_imports(monkeypatch):
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    monkeypatch.setattr(sys, "argv", ["monitor.py", "--list-pos"])
    original = builtins.__import__
    def guarded(name, *values, **kwargs):
        if name in {"config", "data_collector", "telegram_bot", "visualizer", "schedule", "log_masking"}:
            pytest.fail("GHA management request reached an operational import")
        return original(name, *values, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    with pytest.raises(SystemExit) as result:
        runpy.run_path(str(Path(monitor.__file__)), run_name="__main__")
    assert result.value.code == 1


def test_gha_main_never_reinterprets_admin_request_as_normal_job(monkeypatch):
    calls = []
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    monkeypatch.setattr(monitor, "IS_GITHUB_ACTIONS", True)
    monkeypatch.setattr(sys, "argv", ["monitor.py", "--list-pos"])
    monkeypatch.setattr(monitor, "run_github_actions_mode", lambda: calls.append("operational-job"))
    with pytest.raises(SystemExit) as result:
        monitor.main()
    assert result.value.code == 1 and calls == []
